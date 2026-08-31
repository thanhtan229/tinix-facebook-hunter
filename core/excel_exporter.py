"""
Tinix Facebook Down — Excel Exporter
Exports scraped post data to formatted .xlsx files using openpyxl.
"""

import os
import re
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from utils import logger


# Column definitions: (header_text, data_key, width)
COLUMNS = [
    ("STT", "stt", 6),
    ("Ngày post bài", "date", 16),
    ("ID bài viết", "post_id", 22),
    ("Loại bài viết", "post_type", 14),
    ("Engagement", "engagement", 13),
    ("Bình luận", "comments", 12),
    ("Shares", "shares", 13),
    ("Caption", "caption", 60),
    ("Link bài viết", "url", 40),
]

_LINK_FONT = Font(name="Arial", size=10, color="2563EB", underline="single")

# Styling constants
_HEADER_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
_HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
_DATA_FONT = Font(name="Arial", size=10, color="333333")
_CAPTION_FONT = Font(name="Arial", size=10, color="333333")
_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_DATA_ALIGNMENT = Alignment(horizontal="center", vertical="top")
_CAPTION_ALIGNMENT = Alignment(horizontal="left", vertical="top", wrap_text=True)

# Alternating row colors
_ROW_FILL_EVEN = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
_ROW_FILL_ODD = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

# Post type colors
_TYPE_FILLS = {
    "Ảnh": PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"),
    "Video": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    "Text": PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid"),
}


def sanitize_sheet_name(name: str) -> str:
    """Sanitize a string for use as Excel sheet name (max 31 chars, no special chars)."""
    # Remove invalid characters for sheet names
    name = re.sub(r'[\\/*?\[\]:]', '', name)
    # Truncate to 31 characters (Excel limit)
    return name[:31] if name else "Sheet1"


def export_to_excel(
    pages_data: Dict[str, list],
    output_path: str,
) -> bool:
    """
    Export scraped posts to an Excel file.
    Each page gets its own sheet.

    Args:
        pages_data: Dict mapping page_name -> list of ScrapedPost objects.
        output_path: Path to save the .xlsx file.

    Returns:
        True if export was successful.
    """
    try:
        wb = Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        sheet_count = 0
        total_posts = 0

        for page_name, posts in pages_data.items():
            if not posts:
                continue

            sheet_name = sanitize_sheet_name(page_name or f"Page_{sheet_count + 1}")

            # Handle duplicate sheet names
            existing_names = [ws.title for ws in wb.worksheets]
            if sheet_name in existing_names:
                sheet_name = f"{sheet_name}_{sheet_count + 1}"
                sheet_name = sheet_name[:31]

            ws = wb.create_sheet(title=sheet_name)
            sheet_count += 1

            # -- Write header row --
            for col_idx, (header, _, width) in enumerate(COLUMNS, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
                cell.alignment = _HEADER_ALIGNMENT
                cell.border = _BORDER
                ws.column_dimensions[get_column_letter(col_idx)].width = width

            # Set header row height
            ws.row_dimensions[1].height = 32

            # -- Write data rows --
            for row_idx, post in enumerate(posts, 2):
                total_posts += 1
                stt = row_idx - 1

                # Alternating row fill
                row_fill = _ROW_FILL_EVEN if stt % 2 == 0 else _ROW_FILL_ODD

                # Build consistent post URL (numeric ID only)
                numeric_id = post.post_id if post.post_id and post.post_id.isdigit() else ""
                post_url = f"https://www.facebook.com/{numeric_id}" if numeric_id else ""

                # Row data
                row_data = [
                    stt,
                    post.date,
                    post.post_id,
                    post.post_type,
                    post.engagement,
                    post.comments,
                    post.shares,
                    post.caption,
                    post_url,
                ]

                for col_idx, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.font = _DATA_FONT
                    cell.border = _BORDER
                    cell.fill = row_fill

                    # Special styling per column
                    col_key = COLUMNS[col_idx - 1][1]
                    if col_key == "caption":
                        cell.alignment = _CAPTION_ALIGNMENT
                        cell.font = _CAPTION_FONT
                    elif col_key == "post_type":
                        cell.alignment = _DATA_ALIGNMENT
                        type_fill = _TYPE_FILLS.get(value)
                        if type_fill:
                            cell.fill = type_fill
                    elif col_key in ("engagement", "comments", "shares"):
                        cell.alignment = _DATA_ALIGNMENT
                        # Format numbers
                        if isinstance(value, (int, float)) and value > 0:
                            cell.number_format = "#,##0"
                    elif col_key == "url":
                        cell.alignment = Alignment(horizontal="left", vertical="top")
                        cell.font = _LINK_FONT
                        # Add clickable hyperlink
                        if value and isinstance(value, str) and value.startswith("http"):
                            cell.hyperlink = value
                    else:
                        cell.alignment = _DATA_ALIGNMENT

            # Freeze header row
            ws.freeze_panes = "A2"

            # Auto-filter
            if len(posts) > 0:
                ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(posts) + 1}"

        # If no sheets were created, add an empty one
        if sheet_count == 0:
            ws = wb.create_sheet(title="No Data")
            ws.cell(row=1, column=1, value="Không có dữ liệu")
            logger.warning("Không có dữ liệu để export")
            wb.save(output_path)
            return False

        # Save
        output_dir = os.path.dirname(output_path)
        if output_dir:
            os.makedirs(output_dir, exist_ok=True)

        wb.save(output_path)
        logger.info(f"📊 Đã export {total_posts} bài viết ({sheet_count} page) → {os.path.basename(output_path)}")
        return True

    except PermissionError:
        logger.error(f"Không thể ghi file — file đang được mở? ({output_path})")
        return False
    except Exception as e:
        logger.error(f"Lỗi export Excel: {str(e)[:200]}")
        return False
