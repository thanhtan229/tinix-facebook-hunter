"""
Tinix Facebook Down — Scrape Tab
Complete GUI tab for scraping Facebook Page posts.
Saves to SQLite database and auto-exports to Excel.
"""

import os
import random
import threading
import time
from pathlib import Path
from typing import Dict, List, Optional

import customtkinter as ctk

from core.page_scraper import PageScraper, ScrapedPost
from core.db_storage import (
    save_posts, get_db_path, get_db_stats, get_all_groups,
    find_page_by_url, delete_page_data, get_all_accounts,
    get_page_post_ids, get_all_pages_full,
)
from gui.duplicate_dialog import DuplicateDialog, DuplicateAction
from utils import config, logger


class ScrapeTab(ctk.CTkFrame):
    """
    Complete tab for scraping Facebook Page posts → Excel.
    Self-contained: has its own input, controls, progress, and console.
    """

    def __init__(self, master, **kwargs):
        super().__init__(master, fg_color="transparent", **kwargs)

        self._scraper: Optional[PageScraper] = None
        self._running = False
        self._results: Dict[str, List[ScrapedPost]] = {}  # page_name -> posts
        self._page_urls: Dict[str, str] = {}  # page_name -> url
        self._known_post_ids: Dict[str, set] = {}  # url -> set of known post_ids
        self._follower_counts: Dict[str, int] = {}  # page_name -> follower count

        self._setup_ui()

    def _setup_ui(self) -> None:
        """Build the scrape tab layout."""
        # ===== Input Section =====
        input_frame = ctk.CTkFrame(
            self, fg_color="#FFFFFF", corner_radius=12,
            border_width=1, border_color="#E2E8F0",
        )
        input_frame.pack(fill="x", padx=0, pady=(0, 8))

        # URL Input
        url_section = ctk.CTkFrame(input_frame, fg_color="transparent")
        url_section.pack(fill="x", padx=16, pady=(12, 0))

        url_header = ctk.CTkFrame(url_section, fg_color="transparent")
        url_header.pack(fill="x")

        ctk.CTkLabel(
            url_header,
            text="🔗 Facebook Page URLs",
            font=ctk.CTkFont(size=14, weight="bold"),
            text_color="#1E293B",
        ).pack(side="left")

        ctk.CTkButton(
            url_header,
            text="Dán từ Clipboard",
            width=120, height=28,
            font=ctk.CTkFont(size=12),
            fg_color="#EFF6FF", text_color="#2563EB",
            hover_color="#DBEAFE", corner_radius=6,
            command=self._paste_urls,
        ).pack(side="right")

        ctk.CTkLabel(
            url_section,
            text="Nhập URL trang Facebook, mỗi URL một dòng",
            font=ctk.CTkFont(size=11),
            text_color="#94A3B8",
        ).pack(anchor="w", pady=(2, 4))

        self._url_textbox = ctk.CTkTextbox(
            url_section, height=55,
            font=ctk.CTkFont(family="Consolas", size=12),
            fg_color="#F8FAFC", text_color="#334155",
            corner_radius=8, border_width=1, border_color="#E2E8F0",
        )
        self._url_textbox.pack(fill="x")

        # Options row 1: Group + Post count
        row1_frame = ctk.CTkFrame(input_frame, fg_color="transparent")
        row1_frame.pack(fill="x", padx=16, pady=(10, 0))

        # Group selector
        group_frame = ctk.CTkFrame(row1_frame, fg_color="transparent")
        group_frame.pack(side="left")

        ctk.CTkLabel(
            group_frame,
            text="📁 Nhóm:",
            font=ctk.CTkFont(size=13, weight="bold"),
            text_color="#1E293B",
        ).pack(side="left", padx=(0, 8))

        # Load existing groups from DB
        db_path = get_db_path(config.get("scraper_db_path", ""))
        existing_groups = get_all_groups(db_path)

        self._group_var = ctk.StringVar(value="")
        self._group_combo = ctk.CTkComboBox(
            group_frame, width=180, height=36,
            font=ctk.CTkFont(size=13),
            fg_color="#F8FAFC", text_color="#334155",
            border_color="#E2E8F0", corner_radius=8,
            button_color="#CBD5E1", button_hover_color="#94A3B8",
            dropdown_fg_color="#FFFFFF", dropdown_text_color="#1E293B",
            dropdown_hover_color="#EFF6FF",
            values=existing_groups if existing_groups else [""],
            variable=self._group_var,
        )
        self._group_combo.pack(side="left")
        self._group_combo.set("")

        # Post count
        count_frame = ctk.CTkFrame(row1_frame, fg_color="transparent")
        count_frame.pack(side="left", padx=(24, 0))

        ctk.CTkLabel(
            count_frame,
            text="📊 Số bài quét:",
            font=ctk.CTkFont(size=13, weight="bold"),
            text_color="#1E293B",
        ).pack(side="left", padx=(0, 8))

        self._post_count_var = ctk.StringVar(value="ALL")
        self._post_count_entry = ctk.CTkEntry(
            count_frame, width=70, height=36,
            font=ctk.CTkFont(size=13),
            fg_color="#F8FAFC", text_color="#334155",
            border_color="#E2E8F0", corner_radius=8,
            textvariable=self._post_count_var,
            justify="center",
        )
        self._post_count_entry.pack(side="left")

        # Spacer
        ctk.CTkFrame(input_frame, fg_color="transparent", height=8).pack(fill="x", padx=16)

        # ===== Control Section =====
        control_frame = ctk.CTkFrame(
            self, fg_color="#FFFFFF", corner_radius=12,
            border_width=1, border_color="#E2E8F0",
        )
        control_frame.pack(fill="x", padx=0, pady=(0, 8))

        control_inner = ctk.CTkFrame(control_frame, fg_color="transparent")
        control_inner.pack(fill="x", padx=16, pady=10)

        # Left: action buttons
        btn_frame = ctk.CTkFrame(control_inner, fg_color="transparent")
        btn_frame.pack(side="left")

        self._start_btn = ctk.CTkButton(
            btn_frame,
            text="▶  Quét",
            width=120, height=40,
            font=ctk.CTkFont(size=14, weight="bold"),
            fg_color="#2563EB", hover_color="#1D4ED8",
            corner_radius=10,
            command=self._on_start,
        )
        self._start_btn.pack(side="left", padx=(0, 8))

        self._stop_btn = ctk.CTkButton(
            btn_frame,
            text="⏹  Dừng",
            width=100, height=40,
            font=ctk.CTkFont(size=13),
            fg_color="#DC2626", hover_color="#B91C1C",
            corner_radius=10,
            command=self._on_stop,
            state="disabled",
        )
        self._stop_btn.pack(side="left", padx=(0, 8))

        self._results_btn = ctk.CTkButton(
            btn_frame,
            text="📊  Xuất Excel",
            width=120, height=40,
            font=ctk.CTkFont(size=13, weight="bold"),
            fg_color="#16A34A", hover_color="#15803D",
            corner_radius=10,
            command=self._export_results_excel,
            state="disabled",
        )
        self._results_btn.pack(side="left", padx=(0, 8))

        self._pages_report_btn = ctk.CTkButton(
            btn_frame,
            text="📋  Báo cáo Pages",
            width=140, height=40,
            font=ctk.CTkFont(size=13, weight="bold"),
            fg_color="#7C3AED", hover_color="#6D28D9",
            corner_radius=10,
            command=self._export_pages_report,
        )
        self._pages_report_btn.pack(side="left", padx=(0, 8))

        self._reset_btn = ctk.CTkButton(
            btn_frame,
            text="🔄  Reset",
            width=100, height=40,
            font=ctk.CTkFont(size=13),
            fg_color="#F1F5F9", text_color="#475569",
            hover_color="#E2E8F0", corner_radius=10,
            command=self._on_reset,
        )
        self._reset_btn.pack(side="left")



        # ===== Progress Section =====
        progress_frame = ctk.CTkFrame(
            self, fg_color="#FFFFFF", corner_radius=12,
            border_width=1, border_color="#E2E8F0",
        )
        progress_frame.pack(fill="x", padx=0, pady=(0, 8))

        progress_inner = ctk.CTkFrame(progress_frame, fg_color="transparent")
        progress_inner.pack(fill="x", padx=16, pady=12)

        # Status + percentage
        status_row = ctk.CTkFrame(progress_inner, fg_color="transparent")
        status_row.pack(fill="x")

        self._status_label = ctk.CTkLabel(
            status_row,
            text="Sẵn sàng",
            font=ctk.CTkFont(size=13, weight="bold"),
            text_color="#1E293B",
        )
        self._status_label.pack(side="left")

        self._percent_label = ctk.CTkLabel(
            status_row,
            text="",
            font=ctk.CTkFont(size=13, weight="bold"),
            text_color="#2563EB",
        )
        self._percent_label.pack(side="right")

        # Progress bar
        self._progress_bar = ctk.CTkProgressBar(
            progress_inner, height=12, corner_radius=6,
            fg_color="#E2E8F0", progress_color="#2563EB",
        )
        self._progress_bar.pack(fill="x", pady=(6, 8))
        self._progress_bar.set(0)

        # Stats row
        stats_frame = ctk.CTkFrame(progress_inner, fg_color="#F8FAFC", corner_radius=8)
        stats_frame.pack(fill="x")
        stats_inner = ctk.CTkFrame(stats_frame, fg_color="transparent")
        stats_inner.pack(fill="x", padx=12, pady=8)

        self._pages_label = ctk.CTkLabel(
            stats_inner, text="📄 Pages: 0",
            font=ctk.CTkFont(size=12), text_color="#475569",
        )
        self._pages_label.pack(side="left", padx=(0, 24))

        self._posts_label = ctk.CTkLabel(
            stats_inner, text="📝 Bài viết: 0",
            font=ctk.CTkFont(size=12), text_color="#475569",
        )
        self._posts_label.pack(side="left", padx=(0, 24))

        self._time_label = ctk.CTkLabel(
            stats_inner, text="⏱ Thời gian: 00:00",
            font=ctk.CTkFont(size=12), text_color="#475569",
        )
        self._time_label.pack(side="right")

        # ===== Console Section =====
        console_frame = ctk.CTkFrame(
            self, fg_color="#FFFFFF", corner_radius=12,
            border_width=1, border_color="#E2E8F0",
        )
        console_frame.pack(fill="both", expand=True, padx=0, pady=(0, 0))

        console_header = ctk.CTkFrame(console_frame, fg_color="transparent")
        console_header.pack(fill="x", padx=16, pady=(12, 0))

        ctk.CTkLabel(
            console_header,
            text="📋 Console Log",
            font=ctk.CTkFont(size=14, weight="bold"),
            text_color="#1E293B",
        ).pack(side="left")

        ctk.CTkButton(
            console_header,
            text="Xóa Log",
            width=80, height=28,
            font=ctk.CTkFont(size=12),
            fg_color="#F1F5F9", text_color="#64748B",
            hover_color="#E2E8F0", corner_radius=6,
            command=self._clear_console,
        ).pack(side="right")

        self._console = ctk.CTkTextbox(
            console_frame,
            font=ctk.CTkFont(family="Consolas", size=12),
            fg_color="#F8FAFC", text_color="#334155",
            corner_radius=8, border_width=1, border_color="#E2E8F0",
            state="disabled", wrap="word",
        )
        self._console.pack(fill="both", expand=True, padx=12, pady=12)

        # Color tags
        self._console.tag_config("INFO", foreground="#16A34A")
        self._console.tag_config("WARNING", foreground="#D97706")
        self._console.tag_config("ERROR", foreground="#DC2626")
        self._console.tag_config("DEBUG", foreground="#64748B")

    # --- Actions ---

    def _on_start(self) -> None:
        """Start scraping."""
        urls_text = self._url_textbox.get("1.0", "end").strip()
        if not urls_text:
            self._log("ERROR", "Vui lòng nhập ít nhất 1 URL Facebook Page!")
            return

        # Parse post count — "ALL" or empty means unlimited (0)
        count_str = self._post_count_var.get().strip().upper()
        if count_str in ("ALL", ""):
            max_posts = 0  # 0 = unlimited
        else:
            try:
                max_posts = int(count_str)
                if max_posts < 1:
                    raise ValueError
            except (ValueError, TypeError):
                self._log("ERROR", "Số bài quét phải là 'ALL' hoặc số nguyên dương!")
                return

        # Get group name
        group_name = self._group_var.get().strip()

        # Get accounts
        db_path = get_db_path(config.get("scraper_db_path", ""))
        accounts = [a for a in get_all_accounts(db_path) if a["is_active"]]
        if not accounts:
            self._log("ERROR", "Chưa có tài khoản! Vui lòng vào ⚙ Cài đặt → 👤 Quản lý tài khoản")
            return
        acct_names = ', '.join(a.get('label') or a['username'] for a in accounts[:3])
        acct_msg = f"{len(accounts)} tài khoản: {acct_names}"

        # Parse URLs
        urls = self._parse_page_urls(urls_text)
        if not urls:
            self._log("ERROR", "Không tìm thấy URL Facebook Page hợp lệ!")
            return

        # --- Duplicate check (before launching browser) ---
        db_path = get_db_path(config.get("scraper_db_path", ""))
        urls_to_scrape = []
        for url in urls:
            existing = find_page_by_url(db_path, url)
            if existing and existing["post_count"] > 0:
                # Show duplicate dialog
                dlg = DuplicateDialog(
                    self.winfo_toplevel(),
                    page_name=existing["page_name"],
                    page_url=existing["page_url"],
                    post_count=existing["post_count"],
                    last_scraped=existing["last_scraped"],
                    group_name=existing["group_name"],
                )
                action = dlg.get_result()

                if action == DuplicateAction.SKIP:
                    self._log("INFO", f"⏭ Bỏ qua: {url[:60]}")
                    continue
                elif action == DuplicateAction.RESCAN:
                    delete_page_data(db_path, url)
                    self._log("INFO", f"🔄 Xóa dữ liệu cũ, quét lại: {url[:60]}")
                    urls_to_scrape.append(url)
                elif action == DuplicateAction.SCAN_NEW:
                    # Load existing post IDs so scraper knows when to stop
                    existing_ids = get_page_post_ids(db_path, url)
                    self._known_post_ids[url] = existing_ids
                    self._log("INFO", f"🆕 Quét bài mới (bỏ qua {len(existing_ids)} bài đã quét): {url[:60]}")
                    urls_to_scrape.append(url)
                elif action == DuplicateAction.OPEN_FOLDER:
                    output_dir = config.get("output_dir", ".")
                    try:
                        os.startfile(output_dir)
                    except Exception:
                        pass
                    self._log("INFO", f"📂 Mở thư mục: {output_dir}")
                    continue
            else:
                urls_to_scrape.append(url)

        if not urls_to_scrape:
            self._log("WARNING", "Không có URL nào cần quét!")
            return

        # Update UI state
        self._set_running(True)
        self._clear_console()
        self._results = {}
        self._page_urls = {}
        # Note: don't clear _known_post_ids here, it's set before this point
        self._progress_bar.set(0)
        self._percent_label.configure(text="0%")

        limit_text = "ALL" if max_posts == 0 else str(max_posts)
        group_text = f" | Nhóm: {group_name}" if group_name else ""
        self._log("INFO", f"🚀 Bắt đầu quét {len(urls_to_scrape)} page(s), "
                   f"giới hạn: {limit_text} bài{group_text}")
        self._log("INFO", f"👤 {acct_msg} (chọn ngẫu nhiên mỗi URL)")

        # Start in background thread
        thread = threading.Thread(
            target=self._scrape_thread,
            args=(urls_to_scrape, accounts, max_posts, group_name),
            daemon=True,
        )
        thread.start()

    def _on_stop(self) -> None:
        """Stop scraping."""
        if self._scraper:
            self._scraper.request_stop()
        self._log("WARNING", "⏹ Đang dừng...")

    def _export_results_excel(self) -> None:
        """Export scrape results to Excel with full schema columns."""
        if not self._results:
            self._log("WARNING", "Chưa có dữ liệu!")
            return

        try:
            import pandas as pd
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            from datetime import datetime

            all_rows = []
            for page_name, posts in self._results.items():
                page_url = self._page_urls.get(page_name, "")
                group_name = getattr(self, '_current_group', '')
                follower = self._follower_counts.get(page_name, 0)
                for p in posts:
                    pid = p.post_id if p.post_id and p.post_id.isdigit() else ""
                    post_url = f"https://www.facebook.com/{pid}" if pid else ""
                    all_rows.append({
                        "post_id": p.post_id,
                        "page_name": page_name,
                        "page_url": page_url,
                        "group_name": group_name,
                        "date": p.date,
                        "post_type": p.post_type,
                        "engagement": p.engagement,
                        "comments": p.comments,
                        "views": p.views,
                        "caption": p.caption,
                        "url": post_url,
                        "followers": follower,
                    })

            if not all_rows:
                self._log("WARNING", "Không có dữ liệu để xuất!")
                return

            df = pd.DataFrame(all_rows)
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            first_page = list(self._results.keys())[0]
            safe_name = self._sanitize_filename(first_page)
            filename = f"Scrape_{safe_name}_{ts}.xlsx"

            output_dir = config.get("output_dir", ".")
            os.makedirs(output_dir, exist_ok=True)
            filepath = os.path.join(output_dir, filename)

            with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Posts")
                ws = writer.sheets["Posts"]

                hdr_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
                hdr_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")

                for ci in range(1, len(df.columns) + 1):
                    cell = ws.cell(row=1, column=ci)
                    cell.fill = hdr_fill
                    cell.font = hdr_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                    col_letter = get_column_letter(ci)
                    col_name = df.columns[ci - 1]
                    max_len = max(len(str(col_name)), 8)
                    for val in df.iloc[:50, ci - 1]:
                        max_len = max(max_len, min(len(str(val)) if val else 0, 60))
                    ws.column_dimensions[col_letter].width = max_len + 3

                ws.freeze_panes = "A2"
                ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"

            total = len(all_rows)
            self._log("INFO", f"📊 Đã xuất {total} bài viết → {filename}")
            os.startfile(filepath)

        except PermissionError:
            self._log("ERROR", "❌ Không thể ghi file — file Excel đang được mở!")
        except Exception as e:
            self._log("ERROR", f"Lỗi xuất Excel: {str(e)[:200]}")

    def _export_pages_report(self) -> None:
        """Export all scraped pages info to Excel report."""
        try:
            import pandas as pd
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            from datetime import datetime

            db_path = get_db_path(config.get("scraper_db_path", ""))
            pages = get_all_pages_full(db_path)

            if not pages:
                self._log("WARNING", "Database chưa có pages nào!")
                return

            df = pd.DataFrame(pages)
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"Pages_Report_{ts}.xlsx"

            output_dir = config.get("output_dir", ".")
            os.makedirs(output_dir, exist_ok=True)
            filepath = os.path.join(output_dir, filename)

            with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Pages")
                ws = writer.sheets["Pages"]

                hdr_fill = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")
                hdr_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")

                for ci in range(1, len(df.columns) + 1):
                    cell = ws.cell(row=1, column=ci)
                    cell.fill = hdr_fill
                    cell.font = hdr_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                    col_letter = get_column_letter(ci)
                    col_name = df.columns[ci - 1]
                    max_len = max(len(str(col_name)), 10)
                    for val in df.iloc[:50, ci - 1]:
                        max_len = max(max_len, min(len(str(val)) if val else 0, 50))
                    ws.column_dimensions[col_letter].width = max_len + 3

                ws.freeze_panes = "A2"
                ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"

            self._log("INFO", f"📋 Đã xuất báo cáo {len(pages)} pages → {filename}")
            os.startfile(filepath)

        except PermissionError:
            self._log("ERROR", "❌ Không thể ghi file — file Excel đang được mở!")
        except Exception as e:
            self._log("ERROR", f"Lỗi xuất Pages Report: {str(e)[:200]}")

    def _trigger_download(self, urls: list) -> None:
        """Send URLs to the download tab for downloading."""
        try:
            app = self.winfo_toplevel()
            # Switch to download tab
            app._tabview.set("  📥  Tải Media  ")
            # Fill URLs into input panel
            app._input_panel.clear()
            url_text = "\n".join(urls)
            app._input_panel._url_textbox.insert("1.0", url_text)
            self._log("INFO", f"📥 Đã gửi {len(urls)} URL sang tab Tải Media")
            # Auto-start download
            app._on_run()
        except Exception as e:
            self._log("ERROR", f"Lỗi tải: {str(e)[:100]}")

    def _save_to_database(self) -> None:
        """Save all scraped results to SQLite database."""
        if not self._results:
            return

        db_path = get_db_path(config.get("scraper_db_path", ""))
        group_name = getattr(self, '_current_group', '')
        self._log("INFO", f"💾 Đang lưu vào database: {os.path.basename(db_path)}")

        total_saved = 0
        for page_name, posts in self._results.items():
            page_url = self._page_urls.get(page_name, "")
            follower = self._follower_counts.get(page_name, 0)
            saved = save_posts(db_path, page_name, page_url, posts, group_name, follower_count=follower)
            total_saved += saved

        if total_saved > 0:
            stats = get_db_stats(db_path)
            self._log("INFO",
                f"✅ Đã lưu {total_saved} bài | "
                f"Database: {stats['total_posts']} bài từ {stats['total_pages']} page")
        else:
            self._log("WARNING", "⚠ Không lưu được vào database")


    def _on_reset(self) -> None:
        """Reset everything."""
        if self._running:
            self._on_stop()
            time.sleep(0.5)

        self._url_textbox.delete("1.0", "end")
        self._post_count_var.set("ALL")
        self._results = {}
        self._page_urls = {}
        self._known_post_ids = {}
        self._follower_counts = {}
        self._progress_bar.set(0)
        self._percent_label.configure(text="")
        self._status_label.configure(text="Sẵn sàng", text_color="#1E293B")
        self._pages_label.configure(text="📄 Pages: 0")
        self._posts_label.configure(text="📝 Bài viết: 0")
        self._time_label.configure(text="⏱ Thời gian: 00:00")
        self._progress_bar.configure(progress_color="#2563EB")
        self._clear_console()
        self._set_running(False)
        self._results_btn.configure(state="disabled")
        self._log("INFO", "🔄 Đã reset")

    # --- Background Thread ---

    def _scrape_thread(self, urls: List[str], accounts: list, max_posts: int, group_name: str = "") -> None:
        """Run scraping in background thread. Random account per URL."""
        self._current_group = group_name
        start_time = time.time()

        try:
            total_pages = len(urls)
            total_posts = 0

            for i, url in enumerate(urls):
                if self._scraper and self._scraper.is_stopped():
                    break

                # Pick random account for this URL
                account = random.choice(accounts)
                acct_label = account.get('label') or account.get('username', '?')

                self.after(0, self._log, "INFO", f"{'='*40}")
                self.after(0, self._log, "INFO",
                           f"📄 Page {i+1}/{total_pages}: {url[:60]}")
                self.after(0, self._log, "INFO", f"👤 Account: {acct_label}")
                self.after(0, self._update_stats, i, total_pages, total_posts, start_time)

                # Create scraper for this account
                self._scraper = PageScraper(account=account)
                if not self._scraper.start():
                    self.after(0, self._log, "ERROR",
                               f"Không thể khởi tạo Chrome cho {acct_label}!")
                    continue

                try:
                    # Progress callback
                    def on_page_progress(scraped, total, _i=i, _tp=total_pages, _prev=total_posts):
                        frac = (_i / _tp) + (scraped / total / _tp) if total > 0 else 0
                        self.after(0, self._update_progress, frac, _prev + scraped)

                    effective_max = max_posts if max_posts > 0 else 999999
                    
                    # Check if this URL has known post IDs (scan-new mode)
                    url_known_ids = self._known_post_ids.get(url)
                    
                    posts = self._scraper.scrape_page(
                        url, effective_max, on_page_progress,
                        known_post_ids=url_known_ids,
                    )

                    if posts:
                        page_name = posts[0].page_name or f"Page_{i+1}"
                        if page_name in self._results:
                            page_name = f"{page_name} ({i+1})"
                        self._results[page_name] = posts
                        self._page_urls[page_name] = url
                        # Store follower count from scraper
                        self._follower_counts[page_name] = getattr(self._scraper, 'last_follower_count', 0)
                        total_posts += len(posts)
                        self.after(0, self._log, "INFO",
                                   f"✅ '{page_name}': {len(posts)} bài viết")
                    else:
                        if url_known_ids:
                            self.after(0, self._log, "INFO",
                                       f"🆕 Không có bài mới từ: {url[:60]}")
                        else:
                            self.after(0, self._log, "WARNING",
                                       f"⚠ Không quét được bài viết từ: {url[:60]}")
                finally:
                    self._scraper.stop()
                    self._scraper = None

            # Final
            elapsed = time.time() - start_time
            has_scan_new = bool(self._known_post_ids)
            self.after(0, self._scrape_complete, total_posts, elapsed, has_scan_new)

        except Exception as e:
            self.after(0, self._log, "ERROR", f"Lỗi: {str(e)[:200]}")
        finally:
            if self._scraper:
                self._scraper.stop()
                self._scraper = None
            self.after(0, self._set_running, False)

    # --- UI Updates (main thread) ---

    def _scrape_complete(self, total_posts: int, elapsed: float, has_scan_new: bool = False) -> None:
        """Handle scrape completion."""
        mins = int(elapsed // 60)
        secs = int(elapsed % 60)

        self._log("INFO", f"{'='*40}")

        self._pages_label.configure(text=f"📄 Pages: {len(self._results)}")
        self._posts_label.configure(text=f"📝 Bài viết: {total_posts}")
        self._time_label.configure(text=f"⏱ Thời gian: {mins:02d}:{secs:02d}")

        if total_posts > 0:
            # Success
            self._log("INFO", f"✅ HOÀN THÀNH!")
            self._log("INFO", f"   Tổng: {total_posts} bài viết từ {len(self._results)} page(s)")
            self._log("INFO", f"   Thời gian: {mins} phút {secs} giây")

            self._progress_bar.set(1.0)
            self._percent_label.configure(text="100%")
            self._status_label.configure(text="✅ Hoàn thành!", text_color="#16A34A")
            self._progress_bar.configure(progress_color="#16A34A")
            self._results_btn.configure(state="normal")

            # Save to SQLite database
            self._save_to_database()

            # Refresh group dropdown
            self._refresh_groups()

            # Auto-export to Excel
            self._export_results_excel()
        elif has_scan_new:
            # Scan-new mode but no new posts found — this is NOT an error
            self._log("INFO", "🆕 KHÔNG CÓ BÀI MỚI")
            self._log("INFO", "   Page chưa có bài viết mới kể từ lần quét trước")
            self._log("INFO", f"   Thời gian: {mins} phút {secs} giây")

            self._progress_bar.set(1.0)
            self._percent_label.configure(text="100%")
            self._status_label.configure(text="🆕 Không có bài mới", text_color="#7C3AED")
            self._progress_bar.configure(progress_color="#7C3AED")
        else:
            # Failure — no posts scraped
            self._log("ERROR", "❌ THẤT BẠI — Không quét được bài viết nào!")
            self._log("WARNING", "   Có thể do: cookies hết hạn, page private, hoặc URL sai")
            self._log("INFO", f"   Thời gian: {mins} phút {secs} giây")

            self._progress_bar.set(1.0)
            self._percent_label.configure(text="0%")
            self._status_label.configure(text="❌ Thất bại", text_color="#DC2626")
            self._progress_bar.configure(progress_color="#DC2626")

    def _update_progress(self, fraction: float, posts_count: int) -> None:
        """Update progress bar."""
        self._progress_bar.set(min(fraction, 1.0))
        self._percent_label.configure(text=f"{int(fraction * 100)}%")
        self._posts_label.configure(text=f"📝 Bài viết: {posts_count}")

    def _update_stats(self, page_idx: int, total_pages: int, total_posts: int, start_time: float) -> None:
        """Update stats labels."""
        elapsed = time.time() - start_time
        mins = int(elapsed // 60)
        secs = int(elapsed % 60)

        self._status_label.configure(
            text=f"Đang quét page {page_idx + 1}/{total_pages}...",
            text_color="#1E293B",
        )
        self._pages_label.configure(text=f"📄 Pages: {page_idx}/{total_pages}")
        self._posts_label.configure(text=f"📝 Bài viết: {total_posts}")
        self._time_label.configure(text=f"⏱ Thời gian: {mins:02d}:{secs:02d}")

    def _set_running(self, running: bool) -> None:
        """Toggle UI state between running and idle."""
        self._running = running
        if running:
            self._start_btn.configure(state="disabled", fg_color="#93C5FD")
            self._stop_btn.configure(state="normal")
            self._results_btn.configure(state="disabled")
            self._reset_btn.configure(state="disabled")
            self._url_textbox.configure(state="disabled")
            self._post_count_entry.configure(state="disabled")
        else:
            self._start_btn.configure(state="normal", fg_color="#2563EB")
            self._stop_btn.configure(state="disabled")
            self._reset_btn.configure(state="normal")
            self._url_textbox.configure(state="normal")
            self._post_count_entry.configure(state="normal")

    def _log(self, level: str, message: str) -> None:
        """Append a log message to the console."""
        self._console.configure(state="normal")
        tag = level if level in ("INFO", "WARNING", "ERROR", "DEBUG") else "INFO"
        self._console.insert("end", message + "\n", tag)
        self._console.see("end")
        self._console.configure(state="disabled")

    def _clear_console(self) -> None:
        """Clear console."""
        self._console.configure(state="normal")
        self._console.delete("1.0", "end")
        self._console.configure(state="disabled")

    # --- Helpers ---

    def _parse_page_urls(self, text: str) -> List[str]:
        """Parse input text into valid Facebook page URLs."""
        urls = []
        fb_domains = ("facebook.com", "fb.com")

        for line in text.strip().splitlines():
            line = line.strip()
            if not line:
                continue

            # Add protocol if missing
            if not line.startswith(("http://", "https://")):
                line = f"https://{line}"

            try:
                from urllib.parse import urlparse
                parsed = urlparse(line)
                host = parsed.hostname or ""
                if any(host.endswith(d) for d in fb_domains):
                    urls.append(line)
                else:
                    self._log("WARNING", f"Bỏ qua URL không phải Facebook: {line[:60]}")
            except Exception:
                self._log("WARNING", f"URL không hợp lệ: {line[:60]}")

        return urls

    def _paste_urls(self) -> None:
        """Paste clipboard content."""
        try:
            text = self.clipboard_get()
            if text:
                self._url_textbox.insert("end", text)
        except Exception:
            pass



    def _refresh_groups(self) -> None:
        """Refresh group combo dropdown with current DB groups."""
        try:
            db_path = get_db_path(config.get("scraper_db_path", ""))
            groups = get_all_groups(db_path)
            self._group_combo.configure(values=groups if groups else [""])
        except Exception:
            pass

    @staticmethod
    def _sanitize_filename(name: str) -> str:
        """Sanitize for filename use."""
        import re
        name = re.sub(r'[<>:"/\\|?*]', '_', name)
        name = re.sub(r'\s+', '_', name)
        return name.strip('._')[:30] or "page"
