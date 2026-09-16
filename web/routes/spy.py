import os
import re
import mimetypes
from pathlib import Path
from datetime import datetime
from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from pydantic import Field, model_validator
from typing import Annotated, Optional, List, Literal
from core.db_storage import (
    bulk_delete_post_rows,
    get_db_path,
    get_scraped_pages_full,
    get_spy_post_by_row_id,
    query_posts_for_spy,
    MEDIA_DOWNLOAD_EXTS,
    set_post_rating,
    sync_downloaded_posts_from_folder,
)
from core.data_cleanup import cleanup_scraped_data, delete_downloaded_media, preview_scraped_cleanup
from utils import config
from web.ws_manager import log_to_web
from web.models import ApiRequest
from web.routes.downloader import (
    DownloadStartRequest,
    download_has_active_or_pending_work as _download_has_active_or_pending_work,
    download_is_running,
    start_download,
)
from web.routes.scraper import scrape_is_running

router = APIRouter(prefix="/api/spy", tags=["Spy Studio"])

VIDEO_EXTS = {".mp4", ".mov", ".webm", ".avi", ".mkv", ".m4v"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"}


def download_has_active_or_pending_work() -> bool:
    """Keep destructive operations blocked for legacy active work and new queued work."""
    return bool(download_is_running() or _download_has_active_or_pending_work())


class QueryPostsRequest(ApiRequest):
    page_names: List[str] = Field(min_length=1, max_length=1_000)
    mode: Optional[str] = Field(default="newest", pattern="^(newest|oldest|top_engagement|top_reactions|top_comments|top_shares|top_views)$")
    limit: Optional[int] = Field(default=0, ge=0, le=100_000)
    post_type: Optional[str] = Field(default="Tất cả", max_length=100)
    download_status: Optional[str] = Field(default="Tất cả", max_length=100)
    query: Optional[str] = Field(default="", max_length=10_000)
    sort_by: Optional[str] = Field(
        default="", pattern="^(|date|page_name|post_type|engagement|comments|shares|view_count|total_engagement|downloaded)$"
    )
    sort_order: Optional[str] = Field(default="desc", pattern="^(asc|desc)$")


class ExportExcelRequest(ApiRequest):
    page_names: List[str] = Field(min_length=1, max_length=1_000)
    mode: Optional[str] = Field(default="newest", pattern="^(newest|oldest|top_engagement|top_reactions|top_comments|top_shares|top_views)$")
    limit: Optional[int] = Field(default=0, ge=0, le=100_000)
    post_type: Optional[str] = Field(default="Tất cả", max_length=100)
    download_status: Optional[str] = Field(default="Tất cả", max_length=100)
    sort_by: Optional[str] = Field(
        default="", pattern="^(|date|page_name|post_type|engagement|comments|shares|view_count|total_engagement|downloaded)$"
    )
    sort_order: Optional[str] = Field(default="desc", pattern="^(asc|desc)$")


class CleanupScopeRequest(ApiRequest):
    page_urls: List[Annotated[str, Field(min_length=1, max_length=4_096)]] = Field(
        default_factory=list, max_length=1_000
    )
    all_pages: bool = False

    @model_validator(mode="after")
    def validate_scope(self):
        if self.all_pages == bool(self.page_urls):
            raise ValueError("Chọn page_urls hoặc all_pages, không chọn đồng thời")
        return self


class CleanupRequest(CleanupScopeRequest):
    confirmation: Literal["XOA"]
    delete_media: bool = False


class DeletePostsRequest(ApiRequest):
    post_row_ids: List[Annotated[int, Field(gt=0)]] = Field(min_length=1, max_length=100_000)
    delete_media: bool = False


class PostRatingRequest(ApiRequest):
    rating: int = Field(ge=0, le=5)


class RetryPostsRequest(ApiRequest):
    post_row_ids: List[Annotated[int, Field(gt=0)]] = Field(min_length=1, max_length=100_000)
    max_threads: int = Field(default=3, ge=1, le=20)
    use_accounts: Optional[bool] = None
    account_ids: List[Annotated[int, Field(gt=0)]] = Field(default_factory=list, max_length=20)


def _resolve_post_download_path(post: dict, output_root: Path) -> Optional[Path]:
    """Resolve post download folder."""
    raw_path = post.get("download_path")
    if raw_path:
        try:
            cand = Path(raw_path).expanduser().resolve()
            cand.relative_to(output_root)
            if cand.exists():
                return cand
        except (ValueError, OSError):
            pass
    if output_root.is_dir():
        return output_root
    return None


def _downloaded_media_files(post: dict) -> List[Path]:
    """Resolve local media from flat storage without allowing paths outside output_root."""
    output_root = Path(config.get("output_dir", ".")).expanduser().resolve()
    post_id = str(post.get("post_id", "")).strip()
    clean_pid = re.sub(r"\D", "", post_id) or post_id

    if not clean_pid or not output_root.is_dir():
        return []

    try:
        flat_files = [
            p for p in output_root.iterdir()
            if p.is_file() and (
                p.stem == clean_pid
                or p.name.startswith(f"{clean_pid}_")
                or p.name.startswith(f"{clean_pid}.")
            ) and p.suffix.lower() in MEDIA_DOWNLOAD_EXTS
        ]
        if flat_files:
            return sorted(flat_files, key=lambda p: (p.suffix.lower() not in IMAGE_EXTS, p.name.lower()))
    except OSError:
        pass

    raw_path = post.get("download_path")
    if raw_path:
        try:
            cand = Path(raw_path).expanduser().resolve()
            cand.relative_to(output_root)
            if cand.is_dir() and cand != output_root:
                cand_files = [p for p in cand.iterdir() if p.is_file() and p.suffix.lower() in MEDIA_DOWNLOAD_EXTS]
                if cand_files:
                    return sorted(cand_files, key=lambda p: (p.suffix.lower() not in IMAGE_EXTS, p.name.lower()))
            elif cand.is_file() and cand.suffix.lower() in MEDIA_DOWNLOAD_EXTS:
                return [cand]
        except (ValueError, OSError):
            pass

    return []


def _downloaded_caption(post: dict) -> str:
    """Prefer the flat caption file saved with the download, falling back to the database text."""
    output_root = Path(config.get("output_dir", ".")).expanduser().resolve()
    post_id = str(post.get("post_id", "")).strip()
    clean_pid = re.sub(r"\D", "", post_id) or post_id

    if clean_pid and output_root.is_dir():
        flat_caption = output_root / f"{clean_pid}.txt"
        try:
            flat_caption.resolve().relative_to(output_root)
            if flat_caption.is_file() and flat_caption.stat().st_size <= 2_000_000:
                return flat_caption.read_text(encoding="utf-8-sig").strip()
        except (OSError, UnicodeError, ValueError):
            pass

    folder = _resolve_post_download_path(post, output_root)
    if folder and folder != output_root:
        caption_file = folder / "caption.txt" if folder.is_dir() else folder.parent / "caption.txt"
        try:
            caption_file.resolve().relative_to(output_root)
            if caption_file.is_file() and caption_file.stat().st_size <= 2_000_000:
                return caption_file.read_text(encoding="utf-8-sig").strip()
        except (OSError, UnicodeError, ValueError):
            pass
    return post.get("caption", "")


@router.get("/pages")
def get_pages():
    """Get list of all scraped fanpages with metadata."""
    db_p = get_db_path(config.get("scraper_db_path", ""))
    out_dir = config.get("output_dir", "")
    if out_dir:
        sync_downloaded_posts_from_folder(db_p, out_dir)
    pages = get_scraped_pages_full(db_p)
    return {"status": "success", "data": pages}

@router.post("/posts")
def query_posts(req: QueryPostsRequest):
    """Query and filter scraped posts."""
    db_p = get_db_path(config.get("scraper_db_path", ""))
    posts = query_posts_for_spy(
        db_path=db_p,
        page_names=req.page_names,
        mode=req.mode or "newest",
        limit=req.limit or 0,
        post_type_filter=req.post_type or "Tất cả",
        download_status=req.download_status or "Tất cả",
        sort_by=req.sort_by or "",
        sort_order=req.sort_order or "desc",
        query=req.query or "",
    )

    return {
        "status": "success",
        "total": len(posts),
        "data": posts
    }


@router.get("/posts/{post_row_id}/detail")
def get_post_detail(post_row_id: int):
    """Return full text and the complete local media gallery for a downloaded post."""
    db_p = get_db_path(config.get("scraper_db_path", ""))
    post = get_spy_post_by_row_id(db_p, post_row_id)
    if not post:
        raise HTTPException(status_code=404, detail="Không tìm thấy bài viết")
    media_files = _downloaded_media_files(post)
    if not post["downloaded"] and not media_files:
        raise HTTPException(status_code=409, detail="Bài viết chưa được tải về")
    if media_files and not post["downloaded"]:
        post["downloaded"] = True
    post["caption"] = _downloaded_caption(post)
    post.pop("download_path", None)
    post["media"] = [
        {
            "index": index,
            "name": media.name,
            "kind": "video" if media.suffix.lower() in VIDEO_EXTS else "image",
            "url": f"/api/spy/posts/{post_row_id}/media/{index}",
        }
        for index, media in enumerate(media_files)
    ]
    post["has_local_video"] = any(m["kind"] == "video" for m in post["media"])
    return {"status": "success", "data": post}


@router.post("/posts/retry")
def retry_posts(req: RetryPostsRequest):
    """Retry only failed or thumbnail-only Spy posts through the normal downloader."""
    db_p = get_db_path(config.get("scraper_db_path", ""))
    retryable_statuses = {"failed", "thumbnail_only"}
    urls: List[str] = []
    seen_urls: set[str] = set()
    skipped = 0

    for row_id in dict.fromkeys(req.post_row_ids):
        post = get_spy_post_by_row_id(db_p, row_id)
        if not post or post.get("download_status") not in retryable_statuses:
            skipped += 1
            continue
        url = str(post.get("url") or "").strip()
        if not url or url in seen_urls:
            skipped += 1
            continue
        seen_urls.add(url)
        urls.append(url)

    if not urls:
        raise HTTPException(status_code=422, detail="Chỉ có thể thử lại bài Tải lỗi hoặc Thiếu video có link Facebook hợp lệ")

    result = start_download(DownloadStartRequest(
        urls="\n".join(urls),
        max_threads=req.max_threads,
        use_accounts=req.use_accounts,
        account_ids=req.account_ids,
    ))
    result["retrying"] = len(urls)
    result["skipped"] = skipped
    log_to_web(f"Đã đưa {len(urls)} bài lỗi vào lượt tải lại.", "INFO")
    return result


@router.get("/posts/{post_row_id}/media/{media_index}")
def get_post_media(post_row_id: int, media_index: int):
    """Stream one media item selected from the validated downloaded-post folder."""
    db_p = get_db_path(config.get("scraper_db_path", ""))
    post = get_spy_post_by_row_id(db_p, post_row_id)
    if not post:
        raise HTTPException(status_code=404, detail="Không tìm thấy bài viết")
    media_files = _downloaded_media_files(post)
    if media_index < 0 or media_index >= len(media_files):
        raise HTTPException(status_code=404, detail="Không tìm thấy media")
    media = media_files[media_index]
    mime, _ = mimetypes.guess_type(str(media))
    if not mime:
        if media.suffix.lower() in VIDEO_EXTS:
            mime = "video/mp4"
        elif media.suffix.lower() in IMAGE_EXTS:
            mime = "image/jpeg"
    return FileResponse(str(media), media_type=mime or "application/octet-stream")


@router.post("/posts/{post_row_id}/rating")
def rate_post(post_row_id: int, req: PostRatingRequest):
    """Save the user's star rating for one repository post."""
    db_p = get_db_path(config.get("scraper_db_path", ""))
    if not set_post_rating(db_p, post_row_id, req.rating):
        raise HTTPException(status_code=404, detail="Không tìm thấy bài viết")
    return {"status": "success", "data": {"id": post_row_id, "rating": req.rating}}

@router.post("/cleanup-preview")
def cleanup_preview(req: CleanupScopeRequest):
    db_p = get_db_path(config.get("scraper_db_path", ""))
    data = preview_scraped_cleanup(
        db_p, page_urls=req.page_urls, all_pages=req.all_pages
    )
    return {"status": "success", "data": data}


@router.post("/posts/delete")
def delete_posts(req: DeletePostsRequest):
    """Delete exact repository rows and optionally their safely-scoped local media."""
    if scrape_is_running() or download_has_active_or_pending_work():
        raise HTTPException(
            status_code=409,
            detail="Hãy dừng tác vụ quét và tải xuống trước khi xóa bài",
        )
    db_p = get_db_path(config.get("scraper_db_path", ""))
    try:
        result = bulk_delete_post_rows(db_p, req.post_row_ids)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Xóa bài thất bại: {str(exc)[:160]}") from exc
    media_result = {"media_deleted": 0, "media_skipped": 0, "media_failed": 0}
    if req.delete_media:
        media_result = delete_downloaded_media(result["paths"], config.get("output_dir", "."))
    log_to_web(
        f"Đã xóa {result['deleted']} bài khỏi Kho Bài Viết; "
        f"xóa {media_result['media_deleted']} thư mục media.",
        "WARNING",
    )
    return {
        "status": "success",
        "data": {
            "deleted": result["deleted"],
            "media_preserved": 0 if req.delete_media else len(result["paths"]),
            **media_result,
        },
    }


@router.post("/cleanup")
def cleanup_repository(req: CleanupRequest):
    """Back up and remove only scraped pages/posts, optionally scoped media."""
    if scrape_is_running() or download_has_active_or_pending_work():
        raise HTTPException(
            status_code=409,
            detail="Hãy dừng tác vụ quét và tải xuống trước khi dọn kho",
        )
    db_p = get_db_path(config.get("scraper_db_path", ""))
    try:
        result = cleanup_scraped_data(
            db_p,
            page_urls=req.page_urls,
            all_pages=req.all_pages,
            delete_media=req.delete_media,
            output_root=config.get("output_dir", "."),
        )
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Dọn kho thất bại: {str(exc)[:160]}") from exc
    log_to_web(
        f"Đã dọn {result['pages_deleted']} Page và {result['posts_deleted']} bài; "
        f"backup: {result['backup_path'] or 'không cần tạo'}.",
        "WARNING",
    )
    return {"status": "success", "data": result}

@router.post("/export-excel")
def export_excel_endpoint(req: ExportExcelRequest):
    """Export filtered posts to an Excel spreadsheet."""
    import pandas as pd
    from openpyxl.styles import Font, PatternFill, Alignment

    db_p = get_db_path(config.get("scraper_db_path", ""))
    posts = query_posts_for_spy(
        db_path=db_p,
        page_names=req.page_names,
        mode=req.mode or "newest",
        limit=req.limit or 0,
        post_type_filter=req.post_type or "Tất cả",
        download_status=req.download_status or "Tất cả",
        sort_by=req.sort_by or "",
        sort_order=req.sort_order or "desc",
    )

    if not posts:
        raise HTTPException(status_code=400, detail="Không có bài viết nào để xuất Excel!")

    df = pd.DataFrame(posts)
    # Reorder / drop technical columns if present
    drop_cols = [c for c in ["id", "_key"] if c in df.columns]
    if drop_cols:
        df = df.drop(columns=drop_cols)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"Tinix_Posts_Export_{ts}.xlsx"
    out_dir = config.get("output_dir", ".")
    os.makedirs(out_dir, exist_ok=True)
    filepath = os.path.join(out_dir, filename)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Posts")
        ws = writer.sheets["Posts"]
        hdr_fill = PatternFill(start_color="124548", end_color="124548", fill_type="solid")
        hdr_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    log_to_web(f"Đã xuất {len(posts)} bài viết ra file Excel: {filename}", "OK")
    return {
        "status": "success",
        "filename": filename,
        "filepath": os.path.abspath(filepath)
    }
